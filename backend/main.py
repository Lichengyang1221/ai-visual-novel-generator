# -*- coding: utf-8 -*-
"""AI 视觉小说生成器 —— FastAPI 后端。"""
from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import re
import shutil
import time
import uuid
from pathlib import Path

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    datefmt="%H:%M:%S")

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import audio_lib, comfyui_client, config, ollama_client

app = FastAPI(title="AI 视觉小说生成器")


@app.on_event("startup")
async def _auto_resume_on_startup():
    """服务器启动时自动检测有 pending 章节的作品并续跑（防进程死亡后无人恢复）。"""
    import asyncio
    try:
        projects_dir = config.PROJECTS_DIR
        if not projects_dir.exists():
            return
        for d in projects_dir.iterdir():
            mf = d / "manifest.json"
            if not mf.exists():
                continue
            try:
                m = json.loads(mf.read_text(encoding="utf-8"))
            except Exception:
                continue
            pid = d.name
            pending = [ch for ch in m.get("chapters", []) if ch["status"] != "ready"]
            if pending:
                _start_pipeline(pid)
    except Exception:
        pass


@app.middleware("http")
async def no_cache_assets(request, call_next):
    """页面与前端资源不缓存，避免更新后用户看到旧版（图片/音频仍走缓存）。"""
    resp = await call_next(request)
    p = request.url.path
    if p == "/" or p.endswith((".html", ".js", ".css")):
        resp.headers["Cache-Control"] = "no-cache, must-revalidate"
    return resp

# ---------------- 数据模型 ----------------

class StoryIn(BaseModel):
    text: str
    title: str | None = None
    style: str | None = None
    genre: str | None = None


class ScenePatch(BaseModel):
    speaker: str | None = None
    narration: str | None = None
    mood: str | None = None
    scene_prompt_en: str | None = None
    music_track: str | None = None
    anim: str | None = None
    loc_id: str | None = None
    loc_name: str | None = None
    shot: str | None = None
    bg_prompt_en: str | None = None
    lines: list[dict] | None = None  # [{"speaker": str, "text": str}]


class StyleIn(BaseModel):
    style: str | None = None  # None = 跟随全书


class PosIn(BaseModel):
    chapter: int
    scene: int


class BookmarkIn(BaseModel):
    name: str
    chapter: int
    scene: int


class MusicScanIn(BaseModel):
    folder: str


class MusicTrackPatch(BaseModel):
    path: str
    mood: str | None = None  # 大类（兼容旧接口）
    volume: float | None = None  # 0~1
    segments: list[dict] | None = None  # [{"start": 0.35, "end": 0.6, "mood": "静思"}]
    moods: list[str] | None = None  # 多个大类标签
    genres: list[str] | None = None  # 多个作品类型标签


class BatchMusicIn(BaseModel):
    paths: list[str]
    genres: list[str] | None = None  # None=不改；空列表会被规范化为 ["通用"]
    mood: str | None = None  # None=不改；细情绪或大类


# ---------------- 作业（后台任务 + SSE） ----------------

class Job:
    def __init__(self, name: str):
        self.id = uuid.uuid4().hex[:12]
        self.name = name
        self.events: list[dict] = []
        self.done = False
        self.ok = False
        self.queue: asyncio.Queue | None = None

    def emit(self, msg: str, cur: int | None = None, total: int | None = None, **extra):
        ev = {"msg": msg, "cur": cur, "total": total, "t": time.time()}
        ev.update(extra)
        self.events.append(ev)

    def finish(self, ok: bool, msg: str = ""):
        self.done = True
        self.ok = ok
        self.emit(("完成" if ok else "失败") + (f": {msg}" if msg else ""))


JOBS: dict[str, Job] = {}
MUSIC_STATE_FILE = config.APP_ROOT / "music_library.json"

# ---------------- 项目存取 ----------------

def project_dir(pid: str) -> Path:
    d = config.PROJECTS_DIR / pid
    if not d.exists():
        raise HTTPException(404, "项目不存在")
    return d


def migrate(m: dict) -> dict:
    """旧版扁平 scenes manifest 自动迁移为章节结构。"""
    if "chapters" not in m:
        m["chapters"] = [{
            "id": 0, "title": m.get("title", "正文"), "text": "",
            "style": None, "status": "ready",
            "scenes": m.pop("scenes", []),
        }]
    m.setdefault("style", config.DEFAULT_STYLE)
    m.setdefault("genre", config.DEFAULT_GENRE)
    m.setdefault("bookmarks", [])
    m.setdefault("progress", {"chapter": 0, "scene": 0})
    # 书级画面缓存：地点背景图（一个 loc 一张全书复用）、物品 insert 图（按 prompt 复用）
    m.setdefault("locations", {})
    m.setdefault("inserts", {})
    for i, ch in enumerate(m["chapters"]):
        ch.setdefault("id", i)
        ch.setdefault("title", f"第{i+1}章")
        ch.setdefault("text", "")
        ch.setdefault("style", None)
        ch.setdefault("status", "ready")
        ch.setdefault("scenes", [])
        for sc in ch["scenes"]:
            sc.setdefault("anim", None)
            # 三层画面模型字段；旧作品无 loc/shot，播放器按旧模式（每幕一图）播放
            sc.setdefault("loc_id", None)
            sc.setdefault("loc_name", None)
            sc.setdefault("shot", None)  # "close" | "wide" | None(旧数据)
            sc.setdefault("bg_prompt_en", None)
            sc.setdefault("scene_prompt_en", "")
            # 旧作品（一句一幕）迁移为 lines 多句结构
            if not sc.get("lines"):
                sc["lines"] = [{
                    "speaker": sc.get("speaker", "旁白") or "旁白",
                    "text": sc.get("narration", "") or "",
                }]
            # 为旧 lines 补 closeup 字段
            for ln in sc["lines"]:
                ln.setdefault("closeup_prompt_en", None)
                ln.setdefault("closeup_image", None)
            sc.setdefault("narration", "\n".join(ln.get("text", "") for ln in sc["lines"]))
    return m


def load_manifest(pid: str) -> dict:
    m = json.loads((project_dir(pid) / "manifest.json").read_text(encoding="utf-8"))
    return migrate(m)


def save_manifest(pid: str, manifest: dict) -> None:
    (project_dir(pid) / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def get_chapter(m: dict, cid: int) -> dict:
    for ch in m["chapters"]:
        if ch["id"] == cid:
            return ch
    raise HTTPException(404, "章节不存在")


def get_scene(ch: dict, sid: int) -> dict:
    for sc in ch["scenes"]:
        if sc["id"] == sid:
            return sc
    raise HTTPException(404, "场景不存在")


def count_images(m: dict) -> tuple[int, int]:
    """(去重后的实际图片数, 镜头数)。wide 镜头共用 loc 图、insert 跨句复用，故按文件去重。"""
    rels: set[str] = set()
    scenes = 0
    for ch in m["chapters"]:
        for sc in ch["scenes"]:
            scenes += 1
            if sc.get("image_file"):
                rels.add(sc["image_file"])
            for ln in sc.get("lines", []):
                if ln.get("closeup_image"):
                    rels.add(ln["closeup_image"])
    return len(rels), scenes


def load_music_library() -> list[dict]:
    if MUSIC_STATE_FILE.exists():
        try:
            lib = json.loads(MUSIC_STATE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []
    else:
        return []
    # 迁移：早期版本曲目 mood 存的是大类名 → 落成默认细情绪；
    # moods（多选）里的大类也转为该大类默认细情绪。
    changed = False
    for t in lib:
        if t.get("mood") in config.MOOD_CATEGORIES:
            t["mood"] = config.CAT_DEFAULT_MOOD[t["mood"]]
            changed = True
        if isinstance(t.get("moods"), list):
            new_moods = [config.CAT_DEFAULT_MOOD[m] if m in config.MOOD_CATEGORIES else m
                         for m in t["moods"]]
            if new_moods != t["moods"]:
                t["moods"] = new_moods
                changed = True
        # 旧曲目补默认作品类型：通用
        if "genres" not in t:
            t["genres"] = [config.DEFAULT_GENRE]
            changed = True
    if changed:
        save_music_library(lib)
    return lib


def load_deleted_paths() -> set[str]:
    """读取已删除曲目路径集合（防止重新扫描时把删掉的曲子加回来）。"""
    f = config.APP_ROOT / "music_deleted.json"
    if f.exists():
        try:
            return set(json.loads(f.read_text(encoding="utf-8")))
        except Exception:
            return set()
    return set()


def save_deleted_paths(paths: set[str]) -> None:
    f = config.APP_ROOT / "music_deleted.json"
    f.write_text(json.dumps(sorted(paths), ensure_ascii=False, indent=1), encoding="utf-8")


def save_music_library(lib: list[dict]) -> None:
    MUSIC_STATE_FILE.write_text(json.dumps(lib, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------------- 接口：状态与画风 ----------------

@app.get("/api/status")
def api_status():
    return {"ollama": ollama_client.is_running(), "comfyui": comfyui_client.is_running(),
            "moods": config.MOODS}


@app.get("/api/styles")
def api_styles():
    return {"styles": [{"id": k, "name": v["name"]} for k, v in config.STYLES.items()],
            "default": config.DEFAULT_STYLE,
            "anims": config.ANIMS}


# ---------------- 章节生成管线 ----------------

def _img_exists(pid: str, rel: str | None) -> bool:
    return bool(rel) and (project_dir(pid) / rel).exists()


def insert_key(prompt: str) -> str:
    """物品提示词的书级稳定 key（归一化后 md5），同物品跨章复用只出一张图。"""
    norm = ollama_client._pnorm(prompt)
    return hashlib.md5(norm.encode("utf-8")).hexdigest()[:10]


def lens_rel(sc: dict, ch_id: int) -> str:
    """镜头图相对路径：新模式按画面 prompt 的 hash 书级复用（同 prompt 同一张图，
    模型在同一地点重复描写时不重复出图）；旧作品（shot=None）沿用每幕独立文件名。
    L3 前缀：纯黑剪影+意识流模板版本，与旧 L2 图隔离。
    AB 前缀：意识流抽象图，与普通镜头图隔离命名空间。"""
    if sc.get("shot") is None:
        return f"images/c{ch_id:02d}_{sc['id']:03d}.png"
    if sc.get("is_abstract"):
        return f"images/lns_AB_{insert_key(sc['scene_prompt_en'])}.png"
    return f"images/lns_L3_{insert_key(sc['scene_prompt_en'])}.png"


def _gen_lens(pid: str, ch_id: int, sc: dict, style: str, seed: int | None = None) -> str:
    """close 镜头图（特写为主）。is_abstract 场景用意识流模板。"""
    d = project_dir(pid)
    rel = lens_rel(sc, ch_id)
    kind = "abstract" if sc.get("is_abstract") else "lens"
    art_style = load_manifest(pid).get("art_style") or config.DEFAULT_ART_STYLE
    comfyui_client.generate_image(
        sc["scene_prompt_en"], d / rel, seed=seed, style=style, kind=kind,
        art_style=art_style)
    return rel


def _gen_loc(pid: str, loc_id: str, prompt: str, style: str) -> str:
    """地点大背景图（wide establishing，书级复用）。"""
    d = project_dir(pid)
    fname = f"loc_{loc_id}.png"
    art_style = load_manifest(pid).get("art_style") or config.DEFAULT_ART_STYLE
    comfyui_client.generate_image(
        prompt, d / "images" / fname, style=style, kind="bg", art_style=art_style)
    return f"images/{fname}"


def _gen_insert(pid: str, key: str, prompt: str, style: str) -> str:
    """句级物品 insert 图（macro 单品，书级按 prompt 复用）。"""
    d = project_dir(pid)
    fname = f"ins_{key}.png"
    art_style = load_manifest(pid).get("art_style") or config.DEFAULT_ART_STYLE
    comfyui_client.generate_image(
        prompt, d / "images" / fname, style=style, kind="insert",
        art_style=art_style)
    return f"images/{fname}"


# 兼容旧调用名
_gen_one = _gen_lens


async def _pipeline(pid: str, job: Job):
    """逐章生成：缺分镜先拆分镜，再逐张配图，配音乐，就绪一章广播一章。"""
    try:
        lib = [t for t in load_music_library() if "scores" in t]
        manifest = load_manifest(pid)
        total = len(manifest["chapters"])
        # 0) 全书画风基调分析（全书一次，所有 prompt 加基调前缀）
        # 现实主义文学 → realistic，避免 NoobAI 把日常场景画成科幻
        if not manifest.get("art_style"):
            job.emit("全书画风基调分析中…", 0, total, phase="split")
            try:
                sample = "\n".join(ch.get("text", "") for ch in manifest["chapters"][:5])
                art = await asyncio.to_thread(ollama_client.analyze_art_style, sample)
            except Exception:
                art = config.DEFAULT_ART_STYLE
            manifest = load_manifest(pid)
            manifest["art_style"] = art
            save_manifest(pid, manifest)
            job.emit(f"全书画风基调：{config.ART_STYLES.get(art, {}).get('name', art)}",
                     0, total)
        for ci in range(total):
            manifest = load_manifest(pid)
            ch = manifest["chapters"][ci]
            if ch["status"] == "ready":
                continue
            # 1) AI 拆分镜
            if not ch["scenes"]:
                if not ch.get("text"):
                    ch["status"] = "error"
                    save_manifest(pid, manifest)
                    job.emit(f"第 {ci+1} 章缺少原文，跳过", ci, total)
                    continue
                job.emit(f"第 {ci+1}/{total} 章：AI 正在拆分镜…", ci, total,
                         phase="split", split_chars=len(ch["text"]))
                try:
                    res = await asyncio.to_thread(
                        ollama_client.split_story, ch["text"], ch["title"] or None)
                except Exception as e:
                    # 拆分镜失败（含覆盖率多次重试仍不达标）：只标记本章错误，继续后面章节
                    manifest = load_manifest(pid)
                    manifest["chapters"][ci]["status"] = "error"
                    save_manifest(pid, manifest)
                    job.emit(f"第 {ci+1}/{total} 章拆分镜失败，已跳过：{e}", ci, total)
                    continue
                manifest = load_manifest(pid)
                ch = manifest["chapters"][ci]
                ch["scenes"] = res["scenes"]
                if not ch["title"] or ch["title"].startswith("第") and ch["title"].endswith("节"):
                    ch["title"] = ch["title"] or res["title"]
                ch["status"] = "imaging"
                save_manifest(pid, manifest)
            else:
                ch["status"] = "imaging"
                save_manifest(pid, manifest)
            # 单章配图失败不拖垮全书：本章标错继续下一章
            try:
                # 2) 三层配图：①地点大背景（书级复用）②镜头特写 ③物品 insert（书级复用）
                style = ch.get("style") or manifest.get("style") or config.DEFAULT_STYLE
                manifest.setdefault("locations", {})
                manifest.setdefault("inserts", {})

                # 2a) 本章新地点的背景图（同一 loc_id 全书只出一张；文件丢失也会重出）
                todo_loc: list[tuple[str, str]] = []
                loc_seen: set[str] = set()
                for sc in ch["scenes"]:
                    lid = sc.get("loc_id")
                    if not lid or lid in loc_seen:
                        continue
                    loc_seen.add(lid)
                    rec = manifest["locations"].get(lid) or {}
                    prompt = (rec.get("prompt_en")
                              or sc.get("bg_prompt_en") or sc.get("scene_prompt_en"))
                    if prompt and not _img_exists(pid, rec.get("image_file")):
                        todo_loc.append((lid, prompt))
                for k, (lid, prompt) in enumerate(todo_loc):
                    job.emit(f"第 {ci+1}/{total} 章地点背景 {k+1}/{len(todo_loc)}：{lid}…",
                             ci, total, phase="imaging", img_cur=k, img_total=len(todo_loc))
                    rel = await asyncio.to_thread(_gen_loc, pid, lid, prompt, style)
                    manifest = load_manifest(pid)
                    manifest.setdefault("locations", {})[lid] = {
                        "prompt_en": prompt, "image_file": rel}
                    save_manifest(pid, manifest)

                # 2b) wide 建立镜头直接绑定 loc 背景图（不另出）；close 镜头按画面 prompt
                # 的 hash 书级复用：先绑定已存在的图（含跨章复用），缺图再按唯一 prompt 出
                manifest = load_manifest(pid)
                ch = manifest["chapters"][ci]
                for sc in ch["scenes"]:
                    if sc.get("shot") == "wide" and sc.get("loc_id"):
                        rec = manifest["locations"].get(sc["loc_id"]) or {}
                        if _img_exists(pid, rec.get("image_file")):
                            sc["image_file"] = rec["image_file"]
                    elif sc.get("shot") == "close" and sc.get("scene_prompt_en"):
                        rel = lens_rel(sc, ch["id"])
                        if _img_exists(pid, rel):
                            sc["image_file"] = rel
                        else:
                            # 当前版本 hash 的图不存在（模板升级/旧绑定残留）：清掉旧路径，
                            # 否则下面的 todo 过滤会误判"图已存在"而永远不重出
                            sc["image_file"] = None
                save_manifest(pid, manifest)

                todo = [sc for sc in ch["scenes"]
                        if sc.get("shot") != "wide"
                        and not _img_exists(pid, sc.get("image_file"))
                        and (sc.get("scene_prompt_en") or sc.get("shot") is None)]
                # 同 prompt 去重：一章内多个镜头画面相同时只出一张
                uniq: dict[str, dict] = {}
                for sc in todo:
                    uniq.setdefault(lens_rel(sc, ch["id"]), sc)
                for k, (rel, sc) in enumerate(uniq.items()):
                    job.emit(f"第 {ci+1}/{total} 章镜头 {k+1}/{len(uniq)}：{sc['scene_prompt_en'][:36]}…",
                             ci, total, phase="imaging", img_cur=k, img_total=len(uniq))
                    out_rel = await asyncio.to_thread(_gen_lens, pid, ch["id"], sc, style)
                    manifest = load_manifest(pid)
                    # 本章所有同画面镜头统一绑定
                    for bsc in manifest["chapters"][ci]["scenes"]:
                        if bsc.get("shot") == "wide":
                            continue
                        if (bsc.get("scene_prompt_en")
                                and lens_rel(bsc, manifest["chapters"][ci]["id"]) == out_rel):
                            bsc["image_file"] = out_rel
                        elif bsc.get("shot") is None and bsc["id"] == sc["id"]:
                            bsc["image_file"] = out_rel
                    save_manifest(pid, manifest)

                # 2c) 物品 insert：同 prompt（归一化 key）全书只出一张，多处引用共享
                manifest = load_manifest(pid)
                ch = manifest["chapters"][ci]
                ins_map = manifest.setdefault("inserts", {})
                ins_bind: list[tuple[int, int, str]] = []      # (sc_id, li, rel) 直接复用
                ins_todo: dict[str, str] = {}                  # key -> prompt 待生成
                for sc in ch["scenes"]:
                    for li, ln in enumerate(sc.get("lines", [])):
                        cp = ln.get("closeup_prompt_en")
                        if not cp or _img_exists(pid, ln.get("closeup_image")):
                            continue
                        key = insert_key(cp)
                        rec = ins_map.get(key) or {}
                        if _img_exists(pid, rec.get("image_file")):
                            ins_bind.append((sc["id"], li, rec["image_file"]))
                        else:
                            ins_todo.setdefault(key, cp)
                for sc_id, li, rel in ins_bind:
                    get_scene(ch, sc_id)["lines"][li]["closeup_image"] = rel
                if ins_bind:
                    save_manifest(pid, manifest)
                for k, (key, cp) in enumerate(list(ins_todo.items())):
                    job.emit(f"第 {ci+1}/{total} 章物品 {k+1}/{len(ins_todo)}：{cp[:30]}…",
                             ci, total, phase="imaging", img_cur=k, img_total=len(ins_todo))
                    rel = await asyncio.to_thread(_gen_insert, pid, key, cp, style)
                    manifest = load_manifest(pid)
                    manifest.setdefault("inserts", {})[key] = {
                        "prompt_en": cp, "image_file": rel}
                    # 本章所有同 key 行统一绑定（同物品在多句出现）
                    for bsc in manifest["chapters"][ci]["scenes"]:
                        for bln in bsc.get("lines", []):
                            bcp = bln.get("closeup_prompt_en")
                            if bcp and insert_key(bcp) == key and not _img_exists(pid, bln.get("closeup_image")):
                                bln["closeup_image"] = rel
                    save_manifest(pid, manifest)
                # 3) 配乐 + 就绪（聚合配曲：同曲跨镜头沿用，与上一章末尾衔接）
                manifest = load_manifest(pid)
                if lib:
                    prev = None
                    for pch in manifest["chapters"][:ci]:
                        for psc in pch["scenes"]:
                            if psc.get("music_track"):
                                prev = psc["music_track"]
                    seq = manifest["chapters"][ci]["scenes"]
                    tracks = audio_lib.assign_tracks([sc["mood"] for sc in seq], lib,
                                                     start_track=prev, genre=manifest.get("genre", config.DEFAULT_GENRE))
                    for sc, tk in zip(seq, tracks):
                        sc["music_track"] = tk
                manifest["chapters"][ci]["status"] = "ready"
                save_manifest(pid, manifest)
                job.emit(f"第 {ci+1}/{total} 章就绪，可以游玩", ci + 1, total, phase="chapter_ready")
            except Exception as e:
                logger.exception("第 %s 章配图失败：%s", ci + 1, e)
                manifest = load_manifest(pid)
                manifest["chapters"][ci]["status"] = "error"
                save_manifest(pid, manifest)
                job.emit(f"第 {ci+1}/{total} 章配图失败，已跳过：{e}", ci, total)
                continue
        job.finish(True, "全部章节生成完毕")
    except Exception as e:
        try:
            manifest = load_manifest(pid)
            for ch in manifest["chapters"]:
                if ch["status"] in ("pending", "imaging"):
                    ch["status"] = "error"
            save_manifest(pid, manifest)
        except Exception:
            pass
        job.finish(False, str(e))


PIPELINE_PIDS: set[str] = set()


def _start_pipeline(pid: str):
    """同一作品同时只允许一个生成管线（重复点击续跑/多页面操作时直接拒绝）。"""
    if pid in PIPELINE_PIDS:
        return None
    PIPELINE_PIDS.add(pid)
    job = Job("生成章节")
    JOBS[job.id] = job
    task = asyncio.create_task(_pipeline(pid, job))
    task.add_done_callback(lambda _: PIPELINE_PIDS.discard(pid))
    return job


# ---------------- 接口：项目与分镜 ----------------

MAX_UPLOAD_BYTES = 32 * 1024 * 1024  # 文档上传上限 32MB


def _decode_txt(data: bytes) -> str:
    """TXT 编码自动识别：UTF-8(BOM) → UTF-8 → GB18030（中文 Windows 常见 ANSI/GBK）→ Big5 → 兜底。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    """用 python-docx 提取 .docx 正文段落（含表格单元格文字）。"""
    import io

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@app.post("/api/text/extract")
async def api_extract_text(file: UploadFile = File(...)):
    """上传 .txt/.docx，提取纯文本回填到创作框（不直接建项目，便于用户先检查修改）。"""
    name = file.filename or ""
    ext = Path(name).suffix.lower()
    if ext not in (".txt", ".docx"):
        raise HTTPException(400, "仅支持 .txt 和 .docx 文件（旧版 .doc 请先用 Word/WPS 另存为 .docx）")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(400, "文件过大（上限 32MB）")
    if not data:
        raise HTTPException(400, "文件为空")
    try:
        text = await asyncio.to_thread(_extract_docx if ext == ".docx" else _decode_txt, data)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"文档解析失败：{e}（若是旧版 .doc，请先另存为 .docx）")
    # 统一换行、压缩 3 个以上连续空行
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n").replace("\r", "\n")).strip()
    if len(text) < 20:
        raise HTTPException(400, "文档正文太短或为空（至少 20 字）")
    return {"title": Path(name).stem, "text": text, "chars": len(text)}


@app.post("/api/project/create")
async def api_create(story: StoryIn):
    if len(story.text.strip()) < 20:
        raise HTTPException(400, "文本太短（至少 20 字）")
    style = story.style if story.style in config.STYLES else config.DEFAULT_STYLE
    genre = story.genre if story.genre in config.GENRES else config.DEFAULT_GENRE
    try:
        raw_chapters = ollama_client.split_chapters(story.text)
        # 第一章分镜同步生成，保证返回后即可检查
        first = await asyncio.to_thread(
            ollama_client.split_story, raw_chapters[0]["text"], raw_chapters[0]["title"])
    except Exception as e:
        raise HTTPException(500, f"拆分镜失败：{e}")

    pid = time.strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:5]
    d = config.PROJECTS_DIR / pid
    (d / "images").mkdir(parents=True)

    chapters = []
    for i, raw in enumerate(raw_chapters):
        ch = {"id": i, "title": raw["title"], "text": raw["text"],
              "style": None, "status": "pending", "scenes": []}
        if i == 0:
            ch["scenes"] = first["scenes"]
            ch["title"] = ch["title"] or first["title"]
            ch["status"] = "imaging"
        chapters.append(ch)
    manifest = {"id": pid, "title": story.title or first["title"], "created": time.time(),
                "style": style, "genre": genre, "chapters": chapters,
                "progress": {"chapter": 0, "scene": 0}, "bookmarks": []}
    save_manifest(pid, manifest)

    job = _start_pipeline(pid)
    return {"manifest": manifest, "job_id": job.id if job else None}


@app.get("/api/project/list")
def api_list():
    out = []
    for d in sorted(config.PROJECTS_DIR.glob("*"), reverse=True):
        mf = d / "manifest.json"
        if mf.exists():
            try:
                m = migrate(json.loads(mf.read_text(encoding="utf-8")))
            except Exception:
                continue
            imgs, scenes = count_images(m)
            out.append({"id": m["id"], "title": m.get("title", ""),
                        "scenes": scenes, "images": imgs,
                        "chapters": len(m["chapters"]),
                        "ready": sum(1 for c in m["chapters"] if c["status"] == "ready"),
                        "style": m.get("style", config.DEFAULT_STYLE)})
    return out


@app.get("/api/project/{pid}")
def api_get(pid: str):
    return load_manifest(pid)


@app.delete("/api/project/{pid}")
def api_delete(pid: str):
    d = project_dir(pid)
    shutil.rmtree(d)
    return {"deleted": pid}


@app.post("/api/project/{pid}/resume")
async def api_resume(pid: str):
    m = load_manifest(pid)
    if all(ch["status"] == "ready" for ch in m["chapters"]):
        return {"job_id": None, "msg": "全部章节已就绪"}
    job = _start_pipeline(pid)
    if job is None:
        return {"job_id": None, "msg": "该作品正在生成中，请勿重复启动"}
    return {"job_id": job.id}


@app.post("/api/project/{pid}/style")
def api_set_style(pid: str, body: StyleIn):
    if body.style not in config.STYLES:
        raise HTTPException(400, "未知画风")
    m = load_manifest(pid)
    m["style"] = body.style
    save_manifest(pid, m)
    return {"style": body.style}


@app.post("/api/project/{pid}/chapter/{cid}/style")
def api_set_chapter_style(pid: str, cid: int, body: StyleIn):
    if body.style is not None and body.style not in config.STYLES:
        raise HTTPException(400, "未知画风")
    m = load_manifest(pid)
    ch = get_chapter(m, cid)
    ch["style"] = body.style
    save_manifest(pid, m)
    return {"chapter": cid, "style": body.style}


@app.patch("/api/project/{pid}/chapter/{cid}/scene/{sid}")
def api_patch_scene(pid: str, cid: int, sid: int, patch: ScenePatch):
    m = load_manifest(pid)
    sc = get_scene(get_chapter(m, cid), sid)
    if patch.lines is not None:
        old_lines = sc.get("lines", [])
        norm_lines = []
        for li, ln in enumerate(patch.lines):
            text = str(ln.get("text", "")).strip()
            if not text:
                continue
            spk = str(ln.get("speaker", "旁白")).strip() or "旁白"
            entry = {"speaker": spk[:20], "text": text[:300]}
            # 保留 closeup 字段（从前端传入或旧数据对齐）
            cp = str(ln.get("closeup_prompt_en", "") or "").strip()
            entry["closeup_prompt_en"] = cp or None
            # 保留旧的 closeup_image（编辑台词不应丢失已生成的特写图）
            if li < len(old_lines) and old_lines[li].get("closeup_image"):
                entry["closeup_image"] = old_lines[li]["closeup_image"]
            else:
                entry["closeup_image"] = None
            norm_lines.append(entry)
        if not norm_lines:
            raise HTTPException(400, "至少保留一句话")
        sc["lines"] = norm_lines
        sc["narration"] = "\n".join(ln["text"] for ln in norm_lines)
    for k in ("speaker", "narration", "mood", "scene_prompt_en", "music_track", "anim",
              "loc_id", "loc_name", "bg_prompt_en", "shot"):
        v = getattr(patch, k)
        if v is not None:
            if k == "mood" and v not in config.MOODS:
                raise HTTPException(400, f"情绪必须是 {config.MOODS}")
            if k == "anim" and v and v not in config.ANIMS:
                raise HTTPException(400, f"动画必须是 {config.ANIMS} 或空")
            if k == "shot" and v not in ("close", "wide"):
                raise HTTPException(400, "景别只能是 close 或 wide")
            sc[k] = v or None if k == "anim" else v
    save_manifest(pid, m)
    return sc


@app.post("/api/project/{pid}/assign-music")
def api_assign_music(pid: str):
    m = load_manifest(pid)
    lib = load_music_library()
    good = [t for t in lib if "scores" in t]
    if not good:
        raise HTTPException(400, "音乐库为空或尚未分析，请先扫描音乐文件夹")
    # 跨章节按播放顺序集中配曲：一首曲子覆盖连续多个镜头，避免频繁切歌
    seq_scenes = [sc for ch in m["chapters"] for sc in ch["scenes"]]
    tracks = audio_lib.assign_tracks([sc["mood"] for sc in seq_scenes], good,
                                     genre=m.get("genre", config.DEFAULT_GENRE))
    for sc, track in zip(seq_scenes, tracks):
        sc["music_track"] = track
    save_manifest(pid, m)
    return {"assigned": len(seq_scenes), "distinct_tracks": len({t for t in tracks if t})}


@app.post("/api/project/{pid}/chapter/{cid}/scene/{sid}/regen")
async def api_regen(pid: str, cid: int, sid: int):
    """重新出图：wide 建立镜头重出地点背景（全书共享该 loc 的镜头同步换图）；
    close/旧镜头重出该镜头的特写图。"""
    m = load_manifest(pid)
    ch = get_chapter(m, cid)
    sc = get_scene(ch, sid)
    style = ch.get("style") or m.get("style") or config.DEFAULT_STYLE
    try:
        if sc.get("shot") == "wide" and sc.get("loc_id"):
            lid = sc["loc_id"]
            rec = m.setdefault("locations", {}).get(lid) or {}
            prompt = rec.get("prompt_en") or sc.get("bg_prompt_en") or sc.get("scene_prompt_en")
            if not prompt:
                raise HTTPException(400, "该地点缺少背景提示词，无法出图")
            rel = await asyncio.to_thread(_gen_loc, pid, lid, prompt, style)
            m["locations"][lid] = {"prompt_en": prompt, "image_file": rel}
            # 同步全书所有共用该 loc 的 wide 镜头
            for wch in m["chapters"]:
                for wsc in wch["scenes"]:
                    if wsc.get("shot") == "wide" and wsc.get("loc_id") == lid:
                        wsc["image_file"] = rel
        else:
            if not sc.get("scene_prompt_en"):
                raise HTTPException(400, "该镜头缺少画面提示词，无法出图")
            rel = await asyncio.to_thread(_gen_lens, pid, cid, sc, style)
            sc["image_file"] = rel
            # 新模式下同步全书所有同画面（同 prompt）镜头
            if sc.get("shot") is not None:
                for rch in m["chapters"]:
                    for rsc in rch["scenes"]:
                        if rsc.get("shot") == "close" and rsc.get("scene_prompt_en") \
                                and lens_rel(rsc, rch["id"]) == rel:
                            rsc["image_file"] = rel
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"出图失败：{e}")
    save_manifest(pid, m)
    return {"image_file": rel, "t": time.time()}


@app.post("/api/project/{pid}/generate-images")
async def api_generate(pid: str, only_missing: bool = True):
    m = load_manifest(pid)
    job = Job("生成图片")
    JOBS[job.id] = job

    async def run():
        try:
            # ① 地点背景（书级复用）
            manifest = load_manifest(pid)
            todo_loc: dict[str, str] = {}
            for ch0 in manifest["chapters"]:
                style0 = ch0.get("style") or manifest.get("style") or config.DEFAULT_STYLE
                for sc in ch0["scenes"]:
                    lid = sc.get("loc_id")
                    if not lid:
                        continue
                    rec = manifest.setdefault("locations", {}).get(lid) or {}
                    prompt = rec.get("prompt_en") or sc.get("bg_prompt_en") or sc.get("scene_prompt_en")
                    if prompt and (not only_missing or not _img_exists(pid, rec.get("image_file"))):
                        todo_loc.setdefault(lid, prompt)
            total_loc = len(todo_loc)
            for idx, (lid, prompt) in enumerate(todo_loc.items()):
                job.emit(f"地点背景 {idx+1}/{total_loc}：{lid}…", idx, total_loc, phase="genimg")
                rel = await asyncio.to_thread(
                    _gen_loc, pid, lid, prompt,
                    manifest.get("style") or config.DEFAULT_STYLE)
                manifest = load_manifest(pid)
                manifest.setdefault("locations", {})[lid] = {
                    "prompt_en": prompt, "image_file": rel}
                save_manifest(pid, manifest)

            # ② wide 镜头绑定 loc 图；close 镜头先绑定 hash 复用图，再收集缺图（同 prompt 去重）
            manifest = load_manifest(pid)
            targets: dict[str, tuple[int, dict, str]] = {}
            for ch0 in manifest["chapters"]:
                style0 = ch0.get("style") or manifest.get("style") or config.DEFAULT_STYLE
                for sc in ch0["scenes"]:
                    if sc.get("shot") == "wide":
                        if sc.get("loc_id"):
                            rec = manifest["locations"].get(sc["loc_id"]) or {}
                            if _img_exists(pid, rec.get("image_file")):
                                sc["image_file"] = rec["image_file"]
                    elif sc.get("scene_prompt_en") or sc.get("shot") is None:
                        rel = lens_rel(sc, ch0["id"])
                        if _img_exists(pid, rel):
                            sc["image_file"] = rel
                        elif not only_missing or not _img_exists(pid, sc.get("image_file")):
                            targets.setdefault(rel, (ch0["id"], sc, style0))
            save_manifest(pid, manifest)
            total = len(targets)
            for idx, (rel, (cid, sc, style0)) in enumerate(targets.items()):
                job.emit(f"镜头 {idx+1}/{total}：{sc.get('scene_prompt_en', '')[:40]}…",
                         idx, total, phase="genimg")
                out_rel = await asyncio.to_thread(_gen_lens, pid, cid, sc, style0)
                fresh = load_manifest(pid)
                # 全书同画面镜头统一绑定
                for fch in fresh["chapters"]:
                    for fsc in fch["scenes"]:
                        if fsc.get("shot") != "wide" and fsc.get("scene_prompt_en") \
                                and lens_rel(fsc, fch["id"]) == out_rel:
                            fsc["image_file"] = out_rel
                save_manifest(pid, fresh)

            # ③ 物品 insert（书级按 prompt 复用，同 key 全书绑定）
            manifest = load_manifest(pid)
            todo_ins: dict[str, str] = {}
            for ch0 in manifest["chapters"]:
                for sc in ch0["scenes"]:
                    for ln in sc.get("lines", []):
                        cp = ln.get("closeup_prompt_en")
                        if not cp:
                            continue
                        key = insert_key(cp)
                        rec = manifest.setdefault("inserts", {}).get(key) or {}
                        if not only_missing or not _img_exists(pid, rec.get("image_file")):
                            todo_ins.setdefault(key, cp)
                        elif _img_exists(pid, rec.get("image_file")) and not _img_exists(pid, ln.get("closeup_image")):
                            ln["closeup_image"] = rec["image_file"]
            save_manifest(pid, manifest)
            total_ins = len(todo_ins)
            for idx, (key, cp) in enumerate(todo_ins.items()):
                job.emit(f"物品 {idx+1}/{total_ins}：{cp[:40]}…", idx, total_ins, phase="genimg")
                rel = await asyncio.to_thread(
                    _gen_insert, pid, key, cp,
                    manifest.get("style") or config.DEFAULT_STYLE)
                fresh = load_manifest(pid)
                fresh.setdefault("inserts", {})[key] = {"prompt_en": cp, "image_file": rel}
                for bch in fresh["chapters"]:
                    for bsc in bch["scenes"]:
                        for bln in bsc.get("lines", []):
                            bcp = bln.get("closeup_prompt_en")
                            if bcp and insert_key(bcp) == key:
                                bln["closeup_image"] = rel
                save_manifest(pid, fresh)
            job.emit(f"完成：{total_loc} 地点 / {total} 镜头 / {total_ins} 物品",
                     total, total)
            job.finish(True)
        except Exception as e:
            job.finish(False, str(e))

    asyncio.create_task(run())
    return {"job_id": job.id}


# ---------------- 接口：书签与进度 ----------------

@app.get("/api/project/{pid}/bookmarks")
def api_get_bookmarks(pid: str):
    return load_manifest(pid)["bookmarks"]


@app.post("/api/project/{pid}/bookmarks")
def api_add_bookmark(pid: str, bm: BookmarkIn):
    m = load_manifest(pid)
    entry = {"id": uuid.uuid4().hex[:8], "name": bm.name.strip()[:30] or "未命名书签",
             "chapter": bm.chapter, "scene": bm.scene, "t": time.time()}
    m["bookmarks"].append(entry)
    save_manifest(pid, m)
    return entry


@app.delete("/api/project/{pid}/bookmarks/{bid}")
def api_del_bookmark(pid: str, bid: str):
    m = load_manifest(pid)
    before = len(m["bookmarks"])
    m["bookmarks"] = [b for b in m["bookmarks"] if b["id"] != bid]
    if len(m["bookmarks"]) == before:
        raise HTTPException(404, "书签不存在")
    save_manifest(pid, m)
    return {"deleted": bid}


@app.post("/api/project/{pid}/progress")
def api_save_progress(pid: str, pos: PosIn):
    m = load_manifest(pid)
    m["progress"] = {"chapter": pos.chapter, "scene": pos.scene}
    save_manifest(pid, m)
    return m["progress"]


# ---------------- 接口：音乐库 ----------------

@app.post("/api/music/scan")
async def api_music_scan(body: MusicScanIn):
    if not Path(body.folder).is_dir():
        raise HTTPException(400, "文件夹不存在")
    job = Job("分析音乐")
    JOBS[job.id] = job

    async def run():
        def work():
            return audio_lib.analyze_library(
                body.folder, progress_cb=lambda i, n, f: job.emit(f"({i}/{n}) {Path(f).name}", i, n))
        job.emit("开始扫描并分析音乐（首次较慢，之后走缓存）…", 0, 0)
        try:
            lib = await asyncio.to_thread(work)
        except Exception as e:
            job.finish(False, str(e))
            return
        clean = []
        for t in lib:
            if "scores" in t:
                mood = t["mood"]
                if mood in config.MOOD_CATEGORIES:
                    mood = config.CAT_DEFAULT_MOOD[mood]
                clean.append({"path": t["path"], "mood": mood, "scores": t["scores"],
                              "features": t["features"], "genres": [config.DEFAULT_GENRE]})
        # 增量合并：新曲目添加；已有曲目如果 scores 维度不匹配则刷新；跳过已删除的
        existing = load_music_library()
        deleted = load_deleted_paths()
        clean_map = {t["path"]: t for t in clean}
        refreshed = 0
        merged = []
        for t in existing:
            if t["path"] in clean_map and len(t.get("scores", {})) != len(config.MOOD_CATEGORIES):
                # 维度变了，用新分析结果替换，但保留用户已编辑的类型标签
                fresh_t = clean_map[t["path"]]
                if "genres" in t:
                    fresh_t["genres"] = t["genres"]
                merged.append(fresh_t)
                del clean_map[t["path"]]
                refreshed += 1
            else:
                merged.append(t)
        old_paths = {t["path"] for t in merged}
        added = [t for t in clean if t["path"] not in old_paths and t["path"] not in deleted]
        merged += added
        save_music_library(merged)
        skipped = len(lib) - len(clean)
        msg = f"新增 {len(added)} 首"
        if refreshed:
            msg += f"，刷新 {refreshed} 首"
        if skipped:
            msg += f"，跳过 {skipped} 首无法读取"
        msg += f"，曲库共 {len(merged)} 首"
        job.finish(True, msg)

    asyncio.create_task(run())
    return {"job_id": job.id}


@app.get("/api/music/library")
def api_music_library():
    return load_music_library()


@app.delete("/api/music/track")
def api_music_del(path: str):
    """从音乐库中删除一首曲目（仅从库中移除，不删文件，但记住已删除防止重新扫描时恢复）。"""
    lib = load_music_library()
    before = len(lib)
    lib = [t for t in lib if t["path"] != path]
    if len(lib) == before:
        raise HTTPException(404, "曲库中未找到该曲目")
    save_music_library(lib)
    # 记住已删除的路径
    deleted = load_deleted_paths()
    deleted.add(path)
    save_deleted_paths(deleted)
    return {"remaining": len(lib)}


@app.patch("/api/music/track")
def api_music_patch(body: MusicTrackPatch):
    """修改曲目属性：细情绪/音量/播放段落/多情绪标签。大类由细情绪自动归类。"""
    lib = load_music_library()

    def _norm_mood(m: str) -> str | None:
        """接受细情绪或大类，统一返回细情绪。"""
        if m in config.MOODS:
            return m
        if m in config.MOOD_CATEGORIES:
            return config.CAT_DEFAULT_MOOD[m]
        return None

    for t in lib:
        if t["path"] == body.path:
            if body.mood is not None:
                m = _norm_mood(body.mood)
                if m is None:
                    raise HTTPException(400, f"情绪必须为：{'、'.join(config.MOODS)}")
                t["mood"] = m
                cat = config.mood_to_cat(m)
                for c in config.MOOD_CATEGORIES:
                    t["scores"][c] = 0.95 if c == cat else 0.01
            if body.volume is not None:
                t["volume"] = max(0.0, min(1.0, float(body.volume)))
            if body.segments is not None:
                # 段落情绪归一化为细情绪存储
                t["segments"] = []
                for s in body.segments:
                    m = _norm_mood(str(s.get("mood", "平静"))) or "平静"
                    t["segments"].append({"mood": m,
                                          "start": float(s.get("start", 0.35)),
                                          "end": float(s.get("end", 0.6)) if s.get("end") is not None else None})
            if body.moods is not None:
                fine = [m for m in (_norm_mood(x) for x in body.moods) if m]
                # 去重保序
                seen, t["moods"] = set(), []
                for m in fine:
                    if m not in seen:
                        seen.add(m)
                        t["moods"].append(m)
                if t["moods"]:
                    cats = {config.mood_to_cat(m) for m in t["moods"]}
                    for c in config.MOOD_CATEGORIES:
                        t["scores"][c] = 0.9 if c in cats else 0.01
                    t["mood"] = t["moods"][0]
            if body.genres is not None:
                # 作品类型多选：校验、去重保序；空列表视为重置为"通用"
                kept = [g for g in body.genres if g in config.GENRES]
                if not kept:
                    kept = [config.DEFAULT_GENRE]
                seen_g, t["genres"] = set(), []
                for g in kept:
                    if g not in seen_g:
                        seen_g.add(g)
                        t["genres"].append(g)
            save_music_library(lib)
            return {"ok": True, "track": t}
    raise HTTPException(404, "曲库中未找到该曲目")


@app.post("/api/music/batch-edit")
def api_music_batch_edit(body: BatchMusicIn):
    """批量修改曲目的作品类型和/或情绪大类。
    genres=None 表示不改类型；mood=None 表示不改情绪。两者都为 None 报错。"""
    if body.genres is None and body.mood is None:
        raise HTTPException(400, "至少要指定一个要修改的字段（genres 或 mood）")
    if not body.paths:
        raise HTTPException(400, "paths 不能为空")

    lib = load_music_library()
    path_set = set(body.paths)

    # 预处理要应用的值
    apply_genres = None
    if body.genres is not None:
        kept = [g for g in body.genres if g in config.GENRES]
        if not kept:
            kept = [config.DEFAULT_GENRE]
        # 去重保序
        seen, apply_genres = set(), []
        for g in kept:
            if g not in seen:
                seen.add(g)
                apply_genres.append(g)

    apply_mood = None
    if body.mood is not None:
        # 接受细情绪或大类，统一成细情绪
        if body.mood in config.MOODS:
            apply_mood = body.mood
        elif body.mood in config.MOOD_CATEGORIES:
            apply_mood = config.CAT_DEFAULT_MOOD[body.mood]
        else:
            raise HTTPException(400, f"情绪必须为：{'、'.join(config.MOODS)} 或大类：{'、'.join(config.MOOD_CATEGORIES)}")
    apply_cat = config.mood_to_cat(apply_mood) if apply_mood else None

    updated = 0
    for t in lib:
        if t["path"] not in path_set:
            continue
        if apply_genres is not None:
            t["genres"] = list(apply_genres)
        if apply_mood is not None:
            t["mood"] = apply_mood
            t["scores"] = {c: (0.95 if c == apply_cat else 0.01) for c in config.MOOD_CATEGORIES}
        updated += 1
    if updated == 0:
        raise HTTPException(404, "曲库中没有匹配的曲目")
    save_music_library(lib)
    return {"updated": updated, "remaining": len(lib)}


@app.post("/api/music/batch-delete")
def api_music_batch_delete(body: BatchMusicIn):
    """批量从音乐库删除曲目（仅从库移除，不删文件，记住路径防重扫恢复）。"""
    if not body.paths:
        raise HTTPException(400, "paths 不能为空")
    lib = load_music_library()
    path_set = set(body.paths)
    before = len(lib)
    removed_paths = [t["path"] for t in lib if t["path"] in path_set]
    lib = [t for t in lib if t["path"] not in path_set]
    if len(lib) == before:
        raise HTTPException(404, "曲库中没有匹配的曲目")
    save_music_library(lib)
    deleted = load_deleted_paths()
    deleted.update(removed_paths)
    save_deleted_paths(deleted)
    return {"deleted": len(removed_paths), "remaining": len(lib)}


@app.get("/api/music/play-info")
def api_play_info(path: str, mood: str):
    """获取曲目在指定情绪（细情绪或大类）下的播放参数（起播比例、音量、段落）。"""
    lib = load_music_library()
    cat = config.mood_to_cat(mood)
    return audio_lib.get_play_info(path, lib, cat)


@app.get("/audio")
def api_audio(path: str):
    p = Path(path)
    if not p.is_file():
        raise HTTPException(404, "音频不存在")
    return FileResponse(str(p))


# ---------------- 作业 SSE ----------------

@app.get("/api/job/{jid}")
async def api_job(jid: str):
    job = JOBS.get(jid)
    if not job:
        raise HTTPException(404, "作业不存在")

    async def stream():
        sent = 0
        while True:
            while sent < len(job.events):
                ev = job.events[sent]
                sent += 1
                yield f"data: {json.dumps(ev, ensure_ascii=False)}\n\n"
            if job.done:
                yield f"data: {json.dumps({'done': True, 'ok': job.ok}, ensure_ascii=False)}\n\n"
                return
            await asyncio.sleep(0.5)

    return StreamingResponse(stream(), media_type="text/event-stream")


# ---------------- 静态资源 ----------------

app.mount("/projects", StaticFiles(directory=str(config.PROJECTS_DIR)), name="projects")
app.mount("/", StaticFiles(directory=str(config.WEB_DIR), html=True), name="web")
